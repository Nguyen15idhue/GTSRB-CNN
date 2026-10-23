"""Chuyển docs/bao_cao/BAO_CAO.md → BAO_CAO.docx (heading, bảng, hình, caption).

Chạy: python docs/bao_cao/md_to_docx.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parents[2]
MD = ROOT / "docs" / "bao_cao" / "BAO_CAO.md"
DOCX = ROOT / "docs" / "bao_cao" / "BAO_CAO.docx"
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def add_runs(paragraph, text: str) -> None:
    """Thêm text có **bold** và `code` vào paragraph."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(part)


def page_number_footer(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("GTSRB-CNN · Báo cáo bài tập lớn · trang ")
    run = paragraph.add_run()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)


def setup(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.space_after = Pt(4)
    for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
        heading = document.styles[name]
        heading.font.name = "Times New Roman"
        heading.font.size = Pt(size)
        heading.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    section = document.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2)
    page_number_footer(section)


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    n_cols = len(rows[0])
    table = document.add_table(rows=len(rows), cols=n_cols)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(n_cols):
            cell = table.cell(i, j)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            text = row[j] if j < len(row) else ""
            add_runs(paragraph, text)
            for run in paragraph.runs:
                run.font.size = Pt(8.5 if n_cols > 6 else 9.5)
                if i == 0:
                    run.bold = True
    document.add_paragraph()


def main() -> int:
    lines = MD.read_text(encoding="utf-8").splitlines()
    document = Document()
    setup(document)
    index = 0
    first_rule_seen = False
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()

        if not stripped:
            index += 1
            continue

        if stripped == "---":
            if not first_rule_seen:
                document.add_page_break()
                first_rule_seen = True
            index += 1
            continue

        if stripped.startswith("```"):
            block = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            index += 1
            paragraph = document.add_paragraph()
            run = paragraph.add_run("\n".join(block))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            if level == 1:
                paragraph = document.add_heading(text, level=0)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                document.add_heading(text, level=min(level - 1, 3))
            index += 1
            continue

        # bảng
        if stripped.startswith("|"):
            block = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                block.append(lines[index].strip())
                index += 1
            rows = []
            for position, raw in enumerate(block):
                if position == 1 and set(raw.replace("|", "").replace("-", "").strip()) == set():
                    continue
                cells = [cell.strip() for cell in raw.strip("|").split("|")]
                rows.append(cells)
            add_table(document, rows)
            continue

        # hình
        match = re.match(r"!\[(.*?)\]\((.*?)\)", stripped)
        if match:
            caption, path = match.group(1), match.group(2)
            target = ROOT / path
            if target.is_file():
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run().add_picture(str(target), width=Inches(6.2))
            else:
                document.add_paragraph(f"[Thiếu hình: {path}]")
            index += 1
            # caption ở dòng kế tiếp (bỏ qua dòng trống)
            probe = index
            while probe < len(lines) and not lines[probe].strip():
                probe += 1
            if probe < len(lines):
                candidate = lines[probe].strip()
                if candidate.startswith("*") and candidate.endswith("*") and not candidate.startswith("**"):
                    caption_paragraph = document.add_paragraph()
                    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = caption_paragraph.add_run(candidate.strip("*"))
                    run.italic = True
                    run.font.size = Pt(9)
                    index = probe + 1
            continue

        if stripped.startswith("> "):
            text = stripped[2:].strip()
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(1)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(6)
            add_runs(paragraph, text)
            for run in paragraph.runs:
                run.italic = True
            index += 1
            continue

        if stripped.startswith("*") and stripped.endswith("*") and not stripped.startswith("**"):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(stripped.strip("*"))
            run.italic = True
            run.font.size = Pt(9)
            index += 1
            continue

        if re.match(r"^[-*] ", stripped):
            text = stripped[2:]
            index += 1
            while index < len(lines) and lines[index].startswith("  ") and lines[index].strip():
                text += " " + lines[index].strip()
                index += 1
            paragraph = document.add_paragraph(style="List Bullet")
            add_runs(paragraph, text)
            continue

        if re.match(r"^\d+\. ", stripped):
            text = re.sub(r"^\d+\. ", "", stripped)
            index += 1
            while index < len(lines) and lines[index].startswith("   ") and lines[index].strip():
                text += " " + lines[index].strip()
                index += 1
            paragraph = document.add_paragraph(style="List Number")
            add_runs(paragraph, text)
            continue

        # đoạn văn: gộp các dòng liền nhau (Markdown coi xuống dòng đơn là mềm)
        text = stripped
        index += 1
        while index < len(lines):
            candidate = lines[index]
            probe = candidate.strip()
            if not probe:
                break
            if probe.startswith(("#", "|", "![", ">", "```", "---", "*Bảng", "*Hình")):
                break
            if re.match(r"^[-*] ", probe) or re.match(r"^\d+\. ", probe):
                break
            if probe.startswith("*") and probe.endswith("*") and not probe.startswith("**"):
                break
            text += " " + probe
            index += 1
        paragraph = document.add_paragraph()
        add_runs(paragraph, text)

    DOCX.parent.mkdir(parents=True, exist_ok=True)
    document.save(DOCX)
    size_mb = DOCX.stat().st_size / 1024**2
    print(f"Đã ghi {DOCX} ({size_mb:.2f} MB)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
